"""
Scrapes live Torob prices for items that have a torob_url but no torob_1.
Gentle extraction with random jitter to prevent rate limiting.
Updates data/products_data.json, sources/links_torob_and_Digikala.xlsx, and rebuilds dashboard.html.
"""
import json
import logging
import random
import time
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("scrape_missing")

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

from scraper.torob import scrape_torob
from scripts.update_prices import evaluate_torob_top5_status, export_updated_excel
from scripts.build_dashboard_html import generate_dashboard

json_path = BASE_DIR / "data" / "products_data.json"
with open(json_path, "r", encoding="utf-8") as f:
    payload = json.load(f)

items = payload.get("items", [])
missing = [x for x in items if x.get("torob_url") and not x.get("torob_1")]
logger.info("Found %d items missing Torob prices.", len(missing))

excel_path = BASE_DIR / "sources" / "links_torob_and_Digikala.xlsx"
wb = openpyxl.load_workbook(str(excel_path))
ws = wb["Royal_Selected"]

for idx, item in enumerate(missing, 1):
    url = item["torob_url"]
    name = item["name"]
    file_order = item.get("file_order")
    logger.info("[%d/%d] Scraping Torob for (#%s) %s ...", idx, len(missing), file_order, name[:30])
    
    try:
        res = scrape_torob(url)
        if res and res.offers:
            offers = sorted([o for o in res.offers if o.price], key=lambda o: o.price)
            if len(offers) >= 1:
                item["torob_1"] = offers[0].price
                item["torob_price"] = offers[0].price
                item["current_torob_price"] = offers[0].price
            if len(offers) >= 2:
                item["torob_2"] = offers[1].price
            if len(offers) >= 3:
                item["torob_3"] = offers[2].price
            if len(offers) >= 4:
                item["torob_4"] = offers[3].price
            if len(offers) >= 5:
                item["torob_5"] = offers[4].price
            
            # Update excel sheet row
            if file_order:
                row_idx = file_order + 1
                if len(offers) >= 1: ws.cell(row_idx, 15, offers[0].price)
                if len(offers) >= 2: ws.cell(row_idx, 16, offers[1].price)
                if len(offers) >= 3: ws.cell(row_idx, 17, offers[2].price)
                
            logger.info("  -> Got %d offers: [%s, %s, %s, %s, %s]",
                        len(offers),
                        item.get("torob_1"), item.get("torob_2"),
                        item.get("torob_3"), item.get("torob_4"),
                        item.get("torob_5"))
        else:
            logger.info("  -> No offers found or out of stock on Torob.")
    except Exception as e:
        logger.warning("  -> Error: %s", e)

    # Re-evaluate top5 status
    royal_p = item.get("royaldigi_price")
    t1 = item.get("torob_1")
    t2 = item.get("torob_2")
    t3 = item.get("torob_3")
    t4 = item.get("torob_4")
    t5 = item.get("torob_5")
    top5_eval = evaluate_torob_top5_status(royal_p, t1, t2, t3, t4, t5)
    
    item["top5_status"] = top5_eval["status_code"]
    item["top3_status"] = top5_eval["status_code"]
    item["top5_badge"] = top5_eval["badge_text"]
    item["top3_badge"] = top5_eval["badge_text"]
    item["top5_rank_label"] = top5_eval["rank_label"]
    item["top3_rank_label"] = top5_eval["rank_label"]
    item["top5_explanation"] = top5_eval["explanation"]
    item["top3_explanation"] = top5_eval["explanation"]
    item["is_alert"] = top5_eval["is_alert"]
    item["alert_level"] = top5_eval["alert_level"]
    item["diff_from_target"] = top5_eval["diff_from_target"]
    
    # Update sparkline
    if t1 or royal_p:
        item["sparkline"] = [
            royal_p or t1,
            t1 or royal_p,
            t2 or t1 or royal_p,
            t3 or t2 or t1 or royal_p,
            royal_p or t1
        ]
    
    time.sleep(random.uniform(2.0, 3.5))

# Save updated Excel
wb.save(str(excel_path))
logger.info("Saved updated sources/links_torob_and_Digikala.xlsx")

# Update metadata counts
in_top_5_count = sum(1 for x in items if x.get("top5_status") == "in_top_5")
higher_count = sum(1 for x in items if x.get("top5_status") == "higher_than_top_5")
lower_count = sum(1 for x in items if x.get("top5_status") == "lower_than_top_1")
no_royal_count = sum(1 for x in items if x.get("top5_status") == "no_royal")
no_torob_count = sum(1 for x in items if x.get("top5_status") == "no_torob")

payload["metadata"]["in_top_5_count"] = in_top_5_count
payload["metadata"]["in_top_3_count"] = in_top_5_count
payload["metadata"]["higher_than_top_5_count"] = higher_count
payload["metadata"]["higher_than_top_3_count"] = higher_count
payload["metadata"]["lower_than_top_1_count"] = lower_count
payload["metadata"]["no_royal_count"] = no_royal_count
payload["metadata"]["no_torob_count"] = no_torob_count
payload["metadata"]["alerts_total"] = higher_count + lower_count

for out_path in [
    BASE_DIR / "data" / "products_data.json",
    BASE_DIR / "frontend" / "public" / "data" / "products_data.json"
]:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

# Export updated Excel report
excel_report = BASE_DIR / "گزارش_انبار_قیمت_به_روز.xlsx"
export_updated_excel(items, excel_report)
for copy_dest in [
    BASE_DIR / "data" / "گزارش_انبار_قیمت_به_روز.xlsx",
    BASE_DIR / "frontend" / "public" / "data" / "گزارش_انبار_قیمت_به_روز.xlsx"
]:
    import shutil
    shutil.copy2(excel_report, copy_dest)

# Rebuild dashboard.html
generate_dashboard()
logger.info("Successfully finished scraping and rebuilding all dashboards!")
